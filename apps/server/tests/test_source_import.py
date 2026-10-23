"""File/webpage text extraction and the import-script endpoint (3 modes), through the real API."""
import json
from io import BytesIO

import httpx
import pytest
from httpx import ASGITransport, AsyncClient

from src.domains.script import api as script_api
from src.domains.script import series, source_import
from src.domains.script.script_engine import CachedDirectorLLM
from src.main import app
from tests.test_script_engine import claude_response, make_claude

SCRIPT_TWO_EPISODES = "第1集 信\n场1 日 外 天台\n林：你终于来了。\n\n第2集 表\n场1 夜 内 阁楼\n林：表停了。"


def test_extract_file_text_txt_pdf_docx():
    assert source_import.extract_file_text("a.txt", "你好世界".encode()) == "你好世界"
    assert source_import.extract_file_text("a.txt", "你好".encode("gb18030")) == "你好"
    with pytest.raises(ValueError, match="unsupported"):
        source_import.extract_file_text("a.exe", b"x")

    from docx import Document
    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("第一段正文")
    doc.add_paragraph("第二段正文")
    doc.save(buf)
    assert source_import.extract_file_text("a.docx", buf.getvalue()) == "第一段正文\n第二段正文"

    from pypdf import PdfWriter
    buf = BytesIO()
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    w.write(buf)
    assert source_import.extract_file_text("a.pdf", buf.getvalue()) == ""  # blank page, no text


def test_extract_webpage_text_strips_chrome():
    html = ("<html><body><nav>首页 关于 联系</nav><header>Logo</header>"
            "<article><h1>雾港来信第一章</h1><p>林晚在雾港码头等一封信，雨下了整整一夜，没有人知道她在等谁。</p>"
            "<p>另一段讲述了更多的情节细节，用来确认块级文字能被正确保留下来。</p></article>"
            "<footer>版权所有</footer><script>var x=1</script></body></html>")
    text = source_import.extract_webpage_text(html)
    assert "雾港来信第一章" in text and "林晚在雾港码头" in text
    assert "首页" not in text and "版权所有" not in text and "var x=1" not in text


def test_extract_webpage_text_rejects_too_short():
    with pytest.raises(ValueError, match="过短"):
        source_import.extract_webpage_text("<html><body><p>hi</p></body></html>")


async def test_extract_file_endpoint():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        r = await c.post("/api/v1/source/extract-file", files={"file": ("a.txt", b"hello world", "text/plain")})
        assert r.status_code == 200 and r.json() == {"text": "hello world"}
        r = await c.post("/api/v1/source/extract-file", files={"file": ("a.xyz", b"x", "application/octet-stream")})
        assert r.status_code == 422


async def test_fetch_url_endpoint_extracts_body(engine):
    html = "<html><body><article><p>" + "正文内容用来测试抓取网页功能是否可以正常工作。" * 3 + "</p></article></body></html>"

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=html.encode(), headers={"content-type": "text/html; charset=utf-8"})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    from src.main import get_http_client
    app.dependency_overrides[get_http_client] = lambda: client
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
            r = await c.post("/api/v1/source/fetch-url", json={"url": "https://example.test/a"})
            assert r.status_code == 200 and "正文内容" in r.json()["text"]
            r = await c.post("/api/v1/source/fetch-url", json={"url": "not-a-url"})
            assert r.status_code == 422
    finally:
        app.dependency_overrides.clear()


async def test_import_script_as_is_splits_by_episode_marker(engine):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        pid = (await c.post("/api/v1/projects", json={"title": "t"})).json()["id"]
        r = await c.post(f"/api/v1/projects/{pid}/episodes/import",
                         json={"text": SCRIPT_TWO_EPISODES, "mode": "as_is"})
        assert r.status_code == 201 and r.json() == {"episodes": 2}
        eps = (await c.get(f"/api/v1/projects/{pid}/episodes")).json()
        assert [e["status"] for e in eps] == ["scripted", "scripted"]
        assert eps[0]["script_text"].startswith("第1集") and "林：你终于来了" in eps[0]["script_text"]
        assert eps[1]["script_text"].startswith("第2集")
        assert eps[0]["title"] == "信" and eps[1]["title"] == "表"  # marker stripped from the title
        # re-import without replace=true refuses to clobber
        assert (await c.post(f"/api/v1/projects/{pid}/episodes/import",
                             json={"text": SCRIPT_TWO_EPISODES, "mode": "as_is"})).status_code == 409


async def test_import_script_force_as_is_never_splits(engine):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        pid = (await c.post("/api/v1/projects", json={"title": "t"})).json()["id"]
        r = await c.post(f"/api/v1/projects/{pid}/episodes/import",
                         json={"text": SCRIPT_TWO_EPISODES, "mode": "force_as_is"})
        assert r.json() == {"episodes": 1}
        eps = (await c.get(f"/api/v1/projects/{pid}/episodes")).json()
        assert "第2集" in eps[0]["script_text"]  # kept verbatim, including the second marker


async def test_import_script_ai_format_reformats_each_segment(engine, monkeypatch):
    seen = []

    def handler(request):
        body = json.loads(request.content)
        seen.append(body["messages"][0]["content"])
        return claude_response(json.dumps({"episode_title": "x", "script_text": "场1 日 外 X\n甲：整理后的台词。"}))

    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    monkeypatch.setattr(script_api, "CachedDirectorLLM", lambda: CachedDirectorLLM(client=make_claude(handler)))
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as c:
        pid = (await c.post("/api/v1/projects", json={"title": "t"})).json()["id"]
        r = await c.post(f"/api/v1/projects/{pid}/episodes/import",
                         json={"text": SCRIPT_TWO_EPISODES, "mode": "ai_format"})
        assert r.json() == {"episodes": 2}
        eps = (await c.get(f"/api/v1/projects/{pid}/episodes")).json()
        assert all(e["script_text"] == "场1 日 外 X\n甲：整理后的台词。" for e in eps)
        assert len(seen) == 2 and "第1集" in seen[0] and "第2集" in seen[1]


def test_split_episodes():
    assert series.split_episodes("") == []
    assert series.split_episodes("没有标记的一段纯文本") == ["没有标记的一段纯文本"]
    parts = series.split_episodes(SCRIPT_TWO_EPISODES)
    assert len(parts) == 2 and parts[0].startswith("第1集") and parts[1].startswith("第2集")
